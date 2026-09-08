#!/usr/bin/env python3
"""
Build MACRO-format demand CSVs (demand_ref.csv, demand_NDC.csv) from the
provincial hourly load workbook.

Source
  Demand/.../Input_data/regional_hourly_load_profile_2024_analysis.xlsx
    tabs 2021_ref/2025_ref/2030_ref/2035_ref  -> demand_ref.csv
    tabs 2021_NDC/2025_NDC/2030_NDC/2035_NDC  -> demand_NDC.csv
  Each tab: datetime | <national MW> | 17 province MW columns.
  Tabs sit on a 2024 (leap) calendar but hold only 8760 values (Dec 31 empty,
  labels shifted a day after Feb 29). We take the 8760 non-empty rows IN ORDER
  and treat them as Time_Index 1..8760 -- i.e. a plain 365-day hourly series.

Template / schema
  system/demand.csv  -- exact column order and names are reproduced.
  Electricity columns are filled from the workbook; the heat columns
  (Demand_heat*, all flat placeholders) and Demand_Zero are carried over
  unchanged from system/demand.csv.

Output
  system/demand_ref.csv , system/demand_NDC.csv
"""
from __future__ import annotations
import csv
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = (ROOT / "Demand/National_to_Regional_Downscaling_script_py_version"
        "/Input_data/regional_hourly_load_profile_2024_analysis.xlsx")
TEMPLATE = ROOT / "system/demand.csv"
OUT_DIR = ROOT / "system"

NAME2CODE = {
    "Gangwon": "GWN", "Gyeonggi": "GGI", "Gyeongnam": "GNA", "Gyeongbuk": "GNB",
    "Gwangju": "KWJ", "Daegu": "TAE", "Daejeon": "DJJ", "Busan": "PUS",
    "Seoul": "SEL", "Sejong": "SJG", "Ulsan": "USN", "Incheon": "INC",
    "Jeonnam": "JNA", "Jeonbuk": "JNB", "Jeju": "JEJ", "Chungnam": "CNA",
    "Chungbuk": "CNB",
}
YEAR_SUFFIX = {"2021": "", "2025": "_25", "2030": "_30", "2035": "_35"}


def read_tab(ws):
    """Return (national[8760], {code: series[8760]}) from one scenario tab."""
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    # col 0 = datetime, col 1 = national, cols 2.. = provinces (by hdr name)
    prov_col = {}
    for j, h in enumerate(hdr):
        if h in NAME2CODE:
            prov_col[NAME2CODE[h]] = j
    nat, prov = [], {c: [] for c in prov_col}
    for r in rows[1:]:
        if r[1] is None:
            continue
        nat.append(float(r[1]))
        for c, j in prov_col.items():
            prov[c].append(float(r[j]))
    assert len(nat) == 8760, f"{ws.title}: got {len(nat)} rows, expected 8760"
    return nat, prov


def build(scenario: str, wb, template_rows, template_hdr):
    # gather all four year tabs
    data = {}  # suffix -> (nat[8760], {code: series})
    for yr, suf in YEAR_SUFFIX.items():
        data[suf] = read_tab(wb[f"{yr}_{scenario}"])

    out_rows = []
    for i in range(8760):
        src = template_rows[i]           # carry heat / Demand_Zero as-is
        row = dict(src)
        row["Time_Index"] = i + 1
        for suf, (nat, prov) in data.items():
            row[f"Demand_MW{suf}"] = round(nat[i], 6)
            for code, series in prov.items():
                row[f"Demand_MW{suf}_{code}"] = round(series[i], 6)
        out_rows.append(row)

    out = OUT_DIR / f"demand_{scenario}.csv"
    with open(out, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=template_hdr)
        w.writeheader()
        w.writerows(out_rows)
    ann = {suf: sum(nat) / 1e6 for suf, (nat, _) in data.items()}
    pk = {suf: max(nat) for suf, (nat, _) in data.items()}
    print(f"wrote {out.name}")
    for suf, yr in zip(YEAR_SUFFIX.values(), YEAR_SUFFIX):
        print(f"   {yr}: annual {ann[suf]:6.1f} TWh   peak {pk[suf]:8.0f} MW")
    return out


def main():
    with open(TEMPLATE) as fh:
        rd = csv.DictReader(fh)
        template_hdr = rd.fieldnames
        template_rows = list(rd)
    assert len(template_rows) == 8760

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    print("tabs:", wb.sheetnames, "\n")
    build("ref", wb, template_rows, template_hdr)
    print()
    build("NDC", wb, template_rows, template_hdr)


if __name__ == "__main__":
    main()
